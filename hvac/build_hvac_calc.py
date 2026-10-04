"""Builds the SCADA Room HVAC heat-load & unit-selection workbook (formulas live; recalc with LibreOffice)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = 'hvac/SWS-SCADA-Room-HVAC-Heat-Load-and-Selection.xlsx'
F = 'Arial'
BLUE, BLACK, GREEN = '0000FF', '000000', '008000'
HEAD = PatternFill('solid', fgColor='2B2B2B')
KEY = PatternFill('solid', fgColor='FFFF00')
SUB = PatternFill('solid', fgColor='EFEAE0')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)

wb = Workbook()

# ---------------- Design Basis (inputs) ----------------
ws = wb.active; ws.title = 'Design Basis'
ws['A1'] = 'SWS — 40 ft High-Cube E-House · SCADA Room — HVAC Design Basis'; ws['A1'].font = font(bold=True, size=13)
ws['A2'] = 'Legend: blue = input you may change · yellow fill = key assumption to confirm · black = formula · green = link from another sheet'
ws['A2'].font = font(italic=True, size=9)
ws['A3'] = 'Status: FOR REVIEW — design basis assumptions must be confirmed by SWS Engineering before procurement.'
ws['A3'].font = font(color='C00000', bold=True, size=9)
hdr = ['Ref', 'Parameter', 'Value', 'Unit', 'Source / note']
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=5, column=i, value=h); c.font = font('FFFFFF', True); c.fill = HEAD
rows = [
    # ref, label, value, unit, note, key?
    ('—', 'CLIMATE', None, None, None, False),
    ('DB-01', 'Outdoor design dry bulb', 50, '°C', 'Conservative Eastern Province shelter design ambient (ASHRAE 0.4% for Dammam is ≈46 °C). Confirm against project HVAC spec.', True),
    ('DB-02', 'Outdoor design humidity ratio', 22, 'g/kg', 'Coastal humid design condition, Dammam (assumed).', True),
    ('DB-03', 'Indoor design dry bulb', 24, '°C', 'Typical for SCADA / control equipment rooms (24 ± 2 °C).', False),
    ('DB-04', 'Indoor humidity ratio (≈50 % RH @ 24 °C)', 9.3, 'g/kg', 'Psychrometric value at 24 °C / 50 % RH.', False),
    ('—', 'GEOMETRY (from GA drawing SWS-2286 Rev 0)', None, None, None, False),
    ('DB-05', 'External length', 12.192, 'm', 'GA sheet 1: 12,192 mm.', False),
    ('DB-06', 'External width', 2.438, 'm', '40 ft HC container standard width.', False),
    ('DB-07', 'External height', 2.894, 'm', 'GA sheet 1: 2,894 mm.', False),
    ('DB-08', 'Sandwich panel thickness', 0.05, 'm', 'Insulated wall/roof panel (assumed 50 mm PU).', True),
    ('DB-09', 'Internal clear height', 2.55, 'm', 'Assumed after floor build-up and ceiling panel.', True),
    ('—', 'ENVELOPE', None, None, None, False),
    ('DB-10', 'U-value walls & roof', 0.45, 'W/m²K', '50 mm PU sandwich panel incl. surface films (assumed).', True),
    ('DB-11', 'U-value floor', 0.60, 'W/m²K', 'Checker plate on insulated floor, shaded by skid (assumed).', True),
    ('DB-12', 'Solar allowance — roof', 20, 'K', 'Added to ΔT; light-coloured roof (CLTD-simplified).', False),
    ('DB-13', 'Solar allowance — walls', 10, 'K', 'Added to ΔT; light-coloured walls (CLTD-simplified).', False),
    ('DB-14', 'Thermal bridging / door allowance', 0.10, '—', 'Fraction added to envelope gain.', False),
    ('—', 'INTERNAL GAINS', None, None, None, False),
    ('DB-15', 'Equipment heat dissipation (SCADA panels, UPS, DBs)', 6.0, 'kW', 'PLACEHOLDER — replace with electrical heat-dissipation list. Dominant load.', True),
    ('DB-16', 'Lighting power density', 10, 'W/m²', 'LED lighting.', False),
    ('DB-17', 'Occupants (maintenance)', 2, 'persons', 'Intermittent occupancy.', False),
    ('DB-18', 'Sensible heat per person', 75, 'W', 'ASHRAE, seated/light work.', False),
    ('DB-19', 'Latent heat per person', 55, 'W', 'ASHRAE, seated/light work.', False),
    ('—', 'AIR', None, None, None, False),
    ('DB-20', 'Fresh air per person', 10, 'L/s', 'Minimum outdoor air for occupants (assumed).', False),
    ('DB-21', 'Infiltration', 0.3, 'ACH', 'Tight shelter, positive pressure (assumed).', False),
    ('—', 'SELECTION', None, None, None, False),
    ('DB-22', 'Safety factor on design load', 0.10, '—', 'Design margin.', False),
    ('DB-23', 'Number of A/C units installed', 3, 'units', 'As built: 3 split A/C units (project scope).', False),
    ('DB-24', 'Standby units (redundancy)', 1, 'units', 'N+1: one unit on standby / duty rotation.', False),
    ('DB-25', 'Capacity derating at design ambient', 0.80, '—', 'Fraction of nominal (T1) capacity available at design ambient; use T3 (tropical) rated units. Confirm with manufacturer data.', True),
]
ref = {}
r = 6
for rf, label, val, unit, note, key in rows:
    if rf == '—':
        c = ws.cell(row=r, column=2, value=label); c.font = font(bold=True); 
        for col in range(1, 6): ws.cell(row=r, column=col).fill = SUB
    else:
        ws.cell(row=r, column=1, value=rf).font = font(size=9)
        ws.cell(row=r, column=2, value=label).font = font()
        v = ws.cell(row=r, column=3, value=val); v.font = font(BLUE); v.border = BOX
        if key: v.fill = KEY
        if isinstance(val, float) and val < 1 and unit == '—': v.number_format = '0%'
        ws.cell(row=r, column=4, value=unit).font = font(size=9)
        ws.cell(row=r, column=5, value=note).font = font(size=9)
        ref[rf] = f"'Design Basis'!$C${r}"
    r += 1
for col, w in zip('ABCDE', (8, 46, 11, 9, 95)): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A6'

# ---------------- Heat Load ----------------
hl = wb.create_sheet('Heat Load')
hl['A1'] = 'Cooling load calculation — SCADA Room (steady-state, design day)'; hl['A1'].font = font(bold=True, size=13)
hl['A2'] = 'All values in W unless stated. Inputs are linked from the Design Basis sheet (green).'; hl['A2'].font = font(italic=True, size=9)
for i, h in enumerate(['Item', 'Basis', 'Quantity', 'Unit', 'Sensible (W)', 'Latent (W)'], 1):
    c = hl.cell(row=4, column=i, value=h); c.font = font('FFFFFF', True); c.fill = HEAD
R = ref
dT = f"({R['DB-01']}-{R['DB-03']})"
lines = [
    ('Design ΔT (outdoor − indoor)', 'DB-01 − DB-03', f"={dT}", 'K', None, None),
    ('Roof area', 'L × W', f"={R['DB-05']}*{R['DB-06']}", 'm²', None, None),
    ('Wall area (long + end walls)', '2(L+W) × H', f"=2*({R['DB-05']}+{R['DB-06']})*{R['DB-07']}", 'm²', None, None),
    ('Floor area (external)', 'L × W', f"={R['DB-05']}*{R['DB-06']}", 'm²', None, None),
    ('Internal floor area', '(L−2t)(W−2t)', f"=({R['DB-05']}-2*{R['DB-08']})*({R['DB-06']}-2*{R['DB-08']})", 'm²', None, None),
    ('Internal volume', 'internal area × clear height', "=C9*" + R['DB-09'], 'm³', None, None),
    ('Roof conduction + solar', 'U × A × (ΔT + solar roof)', None, None, f"={R['DB-10']}*C6*(C5+{R['DB-12']})", 0),
    ('Wall conduction + solar', 'U × A × (ΔT + solar walls)', None, None, f"={R['DB-10']}*C7*(C5+{R['DB-13']})", 0),
    ('Floor conduction (shaded)', 'U floor × A × ΔT', None, None, f"={R['DB-11']}*C8*C5", 0),
    ('Thermal bridging / doors', 'allowance × envelope', None, None, f"={R['DB-14']}*SUM(E11:E13)", 0),
    ('Equipment heat dissipation', 'DB-15 × 1000', None, None, f"={R['DB-15']}*1000", 0),
    ('Lighting', 'LPD × internal area', None, None, f"={R['DB-16']}*C9", 0),
    ('Occupants', 'n × sensible / latent', None, None, f"={R['DB-17']}*{R['DB-18']}", f"={R['DB-17']}*{R['DB-19']}"),
    ('Fresh air', 'Q = n × L/s; 1.207·Q·ΔT ; 3.0·Q·Δw', f"={R['DB-17']}*{R['DB-20']}", 'L/s', "=1.207*C18*C5", f"=3.0*C18*({R['DB-02']}-{R['DB-04']})"),
    ('Infiltration', 'Q = ACH × V / 3.6', f"={R['DB-21']}*C10/3.6", 'L/s', "=1.207*C19*C5", f"=3.0*C19*({R['DB-02']}-{R['DB-04']})"),
]
r = 5
for item, basis, qty, unit, sens, lat in lines:
    hl.cell(row=r, column=1, value=item).font = font()
    hl.cell(row=r, column=2, value=basis).font = font(size=9)
    if qty is not None:
        c = hl.cell(row=r, column=3, value=qty); c.font = font(GREEN); c.number_format = '#,##0.00'
    if unit: hl.cell(row=r, column=4, value=unit).font = font(size=9)
    for col, v in ((5, sens), (6, lat)):
        if v is not None:
            c = hl.cell(row=r, column=col, value=v); c.font = font(GREEN if isinstance(v, str) and "Design Basis" in v else BLACK); c.number_format = '#,##0'
    r += 1
# r == 20
tot = [
    ('Subtotal', None, '=SUM(E11:E19)', '=SUM(F11:F19)'),
    ('Safety factor', 'DB-22', f"=E20*{R['DB-22']}", f"=F20*{R['DB-22']}"),
    ('DESIGN COOLING LOAD', None, '=E20+E21', '=F20+F21'),
]
for i, (item, basis, s, l) in enumerate(tot):
    rr = 20 + i
    hl.cell(row=rr, column=1, value=item).font = font(bold=True)
    if basis: hl.cell(row=rr, column=2, value=basis).font = font(size=9)
    for col, v in ((5, s), (6, l)):
        c = hl.cell(row=rr, column=col, value=v); c.font = font(bold=True); c.number_format = '#,##0'; c.border = BOX
hl['A24'] = 'Total design load'; hl['A24'].font = font(bold=True)
hl['C24'] = '=(E22+F22)/1000'; hl['C24'].number_format = '0.00'; hl['D24'] = 'kW'
hl['A25'] = 'Total design load'; hl['A25'].font = font(bold=True)
hl['C25'] = '=C24/3.517'; hl['C25'].number_format = '0.00'; hl['D25'] = 'TR'
hl['A26'] = 'Sensible heat ratio'; hl['C26'] = '=IFERROR(E22/(E22+F22),0)'; hl['C26'].number_format = '0.00'
hl['A27'] = 'Equipment share of load'; hl['C27'] = '=IFERROR(E15*(1+' + R['DB-22'] + ')/(E22+F22),0)'; hl['C27'].number_format = '0%'
for c in ('C24', 'C25'): hl[c].font = font(bold=True, size=12); hl[c].fill = SUB
hl['A29'] = 'Method: simplified steady-state CLTD approach (ASHRAE Fundamentals). Air constants: sensible 1.207 W/(L/s·K), latent 3.0 W/(L/s·g/kg).'
hl['A29'].font = font(italic=True, size=9)
for col, w in zip('ABCDEF', (34, 40, 12, 7, 14, 14)): hl.column_dimensions[col].width = w

# ---------------- Unit Selection ----------------
us = wb.create_sheet('Unit Selection')
us['A1'] = 'A/C unit selection — N+1 redundancy check'; us['A1'].font = font(bold=True, size=13)
us['A3'] = 'Design load to cover'; us['C3'] = "='Heat Load'!C24"; us['D3'] = 'kW'
us['A4'] = 'Units installed'; us['C4'] = f"={R['DB-23']}"
us['A5'] = 'Units running (N = installed − standby)'; us['C5'] = f"={R['DB-23']}-{R['DB-24']}"
us['A6'] = 'Derating at design ambient'; us['C6'] = f"={R['DB-25']}"; us['C6'].number_format = '0%'
for c in ('C3', 'C4', 'C5', 'C6'): us[c].font = font(GREEN)
us['C3'].number_format = '0.00'
for i, h in enumerate(['Nominal size (TR)', 'Nominal capacity (kW)', 'Capacity at design ambient (kW)', 'Running units total (kW)', 'Covers load with 1 standby?', 'Capacity margin'], 1):
    c = us.cell(row=8, column=i, value=h); c.font = font('FFFFFF', True); c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center')
sizes = [1.5, 2.0, 2.5, 3.0, 4.0]
for i, s in enumerate(sizes):
    rr = 9 + i
    c = us.cell(row=rr, column=1, value=s); c.font = font(BLUE); c.number_format = '0.0'
    us.cell(row=rr, column=2, value=f'=A{rr}*3.517').number_format = '0.00'
    us.cell(row=rr, column=3, value=f'=B{rr}*$C$6').number_format = '0.00'
    us.cell(row=rr, column=4, value=f'=C{rr}*$C$5').number_format = '0.00'
    us.cell(row=rr, column=5, value=f'=IF(D{rr}>=$C$3,"YES","NO")')
    us.cell(row=rr, column=6, value=f'=IFERROR(D{rr}/$C$3-1,0)').number_format = '0%'
us['A15'] = 'Selected nominal size per unit'; us['A15'].font = font(bold=True)
us['C15'] = '=IFERROR(INDEX(A9:A13,MATCH("YES",E9:E13,0)),"Increase size")'; us['C15'].number_format = '0.0" TR"'
us['A16'] = 'Selected nominal capacity per unit'; us['C16'] = '=IFERROR(C15*3.517,0)'; us['C16'].number_format = '0.00'; us['D16'] = 'kW'
us['A17'] = 'Installed capacity, all units running (at design ambient)'; us['C17'] = '=IFERROR(C16*$C$6*$C$4,0)'; us['C17'].number_format = '0.00'; us['D17'] = 'kW'
us['A18'] = 'Capacity with one unit on standby (at design ambient)'; us['C18'] = '=IFERROR(C16*$C$6*$C$5,0)'; us['C18'].number_format = '0.00'; us['D18'] = 'kW'
for c in ('C15', 'C16', 'C17', 'C18'): us[c].font = font(bold=True); us[c].fill = SUB
us['A20'] = 'Rule: smallest nominal size whose running units (N) cover the design load at design ambient. Use T3 (tropical) rated, R410A/R32 split units with low-ambient and high-ambient protection.'
us['A20'].font = font(italic=True, size=9)
for col, w in zip('ABCDEF', (44, 16, 20, 18, 18, 14)): us.column_dimensions[col].width = w
us.row_dimensions[8].height = 42

# ---------------- Equipment Schedule ----------------
es = wb.create_sheet('Equipment Schedule')
es['A1'] = 'HVAC equipment schedule — SCADA Room'; es['A1'].font = font(bold=True, size=13)
heads = ['Tag', 'Type', 'Location (from left end, m)', 'Nominal size (TR)', 'Nominal capacity (kW)', 'Mounting', 'Notes']
for i, h in enumerate(heads, 1):
    c = es.cell(row=3, column=i, value=h); c.font = font('FFFFFF', True); c.fill = HEAD; c.alignment = Alignment(wrap_text=True)
units = [('AC-01', 2.83), ('AC-02', 6.21), ('AC-03', 9.43)]
r = 4
for tag, x in units:
    for part, mount, note in (('IDU', 'Wall-mounted split indoor unit, high level (≈2.2 m)', 'Lead/lag rotation; F&G trip on fire alarm'),
                              ('ODU', 'External wall bracket, galvanised mesh guard (per GA)', 'T3 rated; condensate & refrigerant via sealed wall sleeve')):
        es.cell(row=r, column=1, value=f'{tag}-{part}').font = font(bold=True)
        es.cell(row=r, column=2, value='Split A/C — ' + ('indoor unit' if part == 'IDU' else 'outdoor condensing unit')).font = font()
        c = es.cell(row=r, column=3, value=x); c.font = font(BLUE); c.number_format = '0.00'
        c = es.cell(row=r, column=4, value="='Unit Selection'!C15"); c.font = font(GREEN); c.number_format = '0.0'
        c = es.cell(row=r, column=5, value="='Unit Selection'!C16"); c.font = font(GREEN); c.number_format = '0.00'
        es.cell(row=r, column=6, value=mount).font = font(size=9)
        es.cell(row=r, column=7, value=note).font = font(size=9)
        r += 1
es.cell(row=r, column=1, value='CTRL-01').font = font(bold=True)
es.cell(row=r, column=2, value='Lead/lag (duty-standby) controller with room thermostat').font = font()
es.cell(row=r, column=6, value='Inside, next to entrance door, 1.5 m AFL').font = font(size=9)
es.cell(row=r, column=7, value='Rotates duty unit every 24 h; starts standby on high temperature (28 °C) or unit fault; alarm to SCADA').font = font(size=9)
es.cell(row=r + 2, column=1, value='Unit positions measured from the A/C outdoor openings on GA sheet 1 elevation (SWS-2286 Rev 0); confirm on site.').font = font(italic=True, size=9)
for col, w in zip('ABCDEFG', (13, 44, 14, 12, 14, 48, 60)): es.column_dimensions[col].width = w
es.row_dimensions[3].height = 32

# ---------------- Notes ----------------
nt = wb.create_sheet('Notes & Assumptions')
notes = [
    'SCOPE: Cooling load and A/C sizing for the 40 ft high-cube SCADA Room E-House (one of three E-Houses in the package).',
    'GEOMETRY from GA drawing SWS-2286 Rev 0 (sheets 1–2), issued 22 Sep 2026.',
    'KEY ASSUMPTIONS (yellow on Design Basis) must be confirmed: outdoor design condition, panel U-values, equipment heat dissipation and unit derating.',
    'EQUIPMENT HEAT DISSIPATION is a placeholder (6.0 kW). It is the dominant load (see Heat Load!C27). Replace with the panel vendor heat-dissipation list.',
    'REDUNDANCY: N+1 — any two units must hold the room at 24 °C at design ambient with one unit failed or on standby.',
    'UNITS: T3 (tropical) rated split units, R410A or R32, epoxy/Blue-fin coated condenser coils for coastal environment, low-ambient kit, auto-restart after power loss.',
    'CONTROLS: lead/lag controller; F&G system trips all A/C units on confirmed fire (per project F&G philosophy); common alarm to SCADA.',
    'CONDENSATE: each IDU drained by gravity (min 1:100 fall) to outside through sealed sleeve, discharged clear of the skid and cable trays.',
    'This workbook is a design-check tool, not a stamped engineering calculation. Status: FOR REVIEW.',
]
nt['A1'] = 'Notes & assumptions'; nt['A1'].font = font(bold=True, size=13)
for i, n in enumerate(notes, 3):
    c = nt.cell(row=i, column=1, value=f'{i-2}. {n}'); c.font = font(); c.alignment = Alignment(wrap_text=True, vertical='top')
nt.column_dimensions['A'].width = 140

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
wb.save(OUT); print('saved', OUT)
